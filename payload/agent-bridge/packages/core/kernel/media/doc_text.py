"""把对方发来的文件读成纯文字（给模型看）：txt/代码/字幕等文本、PDF、Word、Excel、PPT。

python media/doc_text.py <文件> [最多多少字，默认 4000]
输出 JSON：{"text": "...", "truncated": true/false, "kind": "pdf"}；读不了 text 为空、带 error。
"""
from __future__ import annotations

import json
import sys
from pathlib import Path


def read_text(p: Path) -> str:
    raw = p.read_bytes()
    for enc in ("utf-8-sig", "gb18030", "utf-16"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def read_pdf(p: Path, limit: int) -> str:
    import fitz  # PyMuPDF

    out = []
    with fitz.open(p) as doc:
        for page in doc:
            out.append(page.get_text())
            if sum(map(len, out)) > limit:
                break
    return "\n".join(out)


def read_docx(p: Path, limit: int) -> str:
    import docx

    d = docx.Document(str(p))
    lines = [x.text for x in d.paragraphs if x.text.strip()]
    for t in d.tables:
        for row in t.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
            if sum(map(len, lines)) > limit:
                break
    return "\n".join(lines)


def read_xlsx(p: Path, limit: int) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets[:3]:
        out.append(f"【{ws.title}】")
        for row in ws.iter_rows(max_row=200, values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(cells):
                out.append(" | ".join(cells).rstrip(" |"))
            if sum(map(len, out)) > limit:
                break
    return "\n".join(out)


def read_pptx(p: Path, limit: int) -> str:
    from pptx import Presentation

    out = []
    for i, slide in enumerate(Presentation(str(p)).slides, 1):
        texts = [sh.text_frame.text for sh in slide.shapes if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip()]
        if texts:
            out.append(f"【第 {i} 页】" + " / ".join(texts))
        if sum(map(len, out)) > limit:
            break
    return "\n".join(out)


def main() -> None:
    p = Path(sys.argv[1])
    limit = int(sys.argv[2]) if len(sys.argv) > 2 else 4000
    ext = p.suffix.lower().lstrip(".")
    try:
        if ext == "pdf":
            text = read_pdf(p, limit)
        elif ext == "docx":
            text = read_docx(p, limit)
        elif ext == "xlsx":
            text = read_xlsx(p, limit)
        elif ext == "pptx":
            text = read_pptx(p, limit)
        else:
            text = read_text(p)
        text = text.replace("\r\n", "\n").strip()
        print(json.dumps({"text": text[:limit], "truncated": len(text) > limit, "kind": ext}, ensure_ascii=False))
    except Exception as e:  # 读不了：交给模型自己想办法
        print(json.dumps({"text": "", "truncated": False, "kind": ext, "error": f"{type(e).__name__}: {e}"}, ensure_ascii=False))


if __name__ == "__main__":
    main()
